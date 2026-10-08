from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT='Quick_GIGS_Business_Plan_and_GTM.docx'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)

# palette
INK='101014'; MUTED='5F6368'; MINT='00A878'; LIGHT='EAF8F3'; LINE='D9E2DE'; LAV='F4F0FF'
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string(INK)
styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size,color in [('Title',34,INK),('Heading 1',21,INK),('Heading 2',14,INK),('Heading 3',10.5,MINT)]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(5)
# custom small label
if 'Eyebrow' not in styles:
    st=styles.add_style('Eyebrow', WD_STYLE_TYPE.PARAGRAPH); st.font.name='Aptos'; st.font.size=Pt(8); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(MINT); st.font.all_caps=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(cell, color=LINE):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),color); b.append(x)
    tcPr.append(b)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,INK); borders(c,INK); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8.5)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); borders(cells[i]);
            if ri%2==0: shade(cells[i],'F7FAF8')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(8.3)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p
def para(text='', boldlead=None):
    p=doc.add_paragraph()
    if boldlead and text.startswith(boldlead): p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def page(): doc.add_page_break()

# header/footer
header=sec.header.paragraphs[0]; header.text='QUICK GIGS  /  BUSINESS PLAN + GO-TO-MARKET'; header.style='Eyebrow'
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Confidential working plan  •  September 2026  •  Vinayak Tower, Vibhuti Khand, Lucknow 226028').font.size=Pt(8)

# cover
p=doc.add_paragraph('BUSINESS PLAN',style='Eyebrow'); p.paragraph_format.space_before=Pt(50)
p=doc.add_paragraph('Quick GIGS',style='Title'); p.runs[0].font.color.rgb=RGBColor.from_string(MINT)
p=doc.add_paragraph('The fastest, most transparent way for Indian businesses to get creative work made.'); p.runs[0].font.size=Pt(18); p.runs[0].font.bold=True
p=doc.add_paragraph('Business model, positioning, launch plan and 12-month go-to-market strategy'); p.runs[0].font.size=Pt(12); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
doc.add_paragraph('\n')
t=doc.add_table(rows=4, cols=2); t.alignment=WD_TABLE_ALIGNMENT.LEFT
for i,(a,b) in enumerate([('Business','Quick GIGS'),('Base','Lucknow, India'),('Address','Vinayak Tower, Vibhuti Khand, Lucknow 226028'),('Planning horizon','12 months from launch')] ):
    t.cell(i,0).text=a; t.cell(i,1).text=b; shade(t.cell(i,0),LIGHT); borders(t.cell(i,0)); borders(t.cell(i,1))
    t.cell(i,0).paragraphs[0].runs[0].font.bold=True
para('Planning note: This is a practical operating plan built from the current product direction. Market sizes, conversion rates and financial outcomes are planning assumptions—not audited forecasts—and should be replaced with live data after the first 30–60 days.')
page()

# executive
h=doc.add_heading('1. Executive summary',1)
para('Quick GIGS is a managed, AI-assisted marketplace for creative and digital services. A business describes the work in plain language; the platform turns it into a production-ready brief, matches a verified freelancer or creator, holds customer funds in escrow, and coordinates approval and payout.')
para('The wedge is not “another freelancer directory.” It is speed plus trust: less time writing briefs, less time searching, clearer pricing, and a safer delivery process. The initial focus should be Indian SMBs, D2C brands, agencies and creator-led companies that need recurring content but do not want a full-time creative team.')
# callout
c=doc.add_table(rows=1,cols=1).cell(0,0); shade(c,LIGHT); borders(c,MINT); c.text='North-star promise: “Describe the work. Get a verified specialist and a clear next step—without weeks of waiting.”'
c.paragraphs[0].runs[0].bold=True

doc.add_heading('12-month objectives',2)
table(['Objective','Target / decision rule'],[
('Demand','Reach 1,000 qualified business leads; 250 completed first orders.'),
('Supply','Build 150 active, verified specialists and 20 agency/team profiles.'),
('Repeat','Achieve 35% second-order rate within 90 days of first order.'),
('Quality','Maintain ≥90% on-time delivery and ≥4.6/5 post-delivery rating.'),
('Economics','Reach contribution-positive transactions before scaling paid acquisition.'),
('Trust','Keep creator fee at 0%; show only service subtotal + GST to customers.')])

# problem solution
h=doc.add_heading('2. Problem, customer and solution',1)
doc.add_heading('Customer pain',2)
for x in ['SMBs need reels, ads, thumbnails, landing pages, copy and UGC every week—but hiring is slow and fragmented.','Briefs are vague, so revisions multiply and the buyer cannot tell whether the price is fair.','Freelancers lose time chasing leads and fear delayed or unclear payouts.','Agencies need overflow capacity and specialist pods without exposing their client relationship.']: bullet(x)
doc.add_heading('Product solution',2)
table(['Stage','Customer outcome','Platform capability'],[
('1. Describe','A rough idea becomes a structured request.','AI brief writer, category and scope prompts.'),
('2. Match','Shortlist based on skill, availability, history and budget.','Explainable matching and verified profiles.'),
('3. Order','Know exactly what is included and what GST is.','Service pricing, GST line item, escrow.'),
('4. Deliver','See progress and submit feedback without chaos.','Milestones, chat, delivery and QA gate.'),
('5. Approve','Release funds only when work is accepted.','Approval workflow and freelancer payout.')])

# market
h=doc.add_heading('3. Market focus and positioning',1)
doc.add_heading('Beachhead segments',2)
table(['Segment','Trigger','Initial offer','Why now'],[
('D2C / consumer brands','Weekly product launches and paid social.','“10 reels + thumbnails in 30 days.”','Content velocity directly affects ad testing.'),
('Local and regional SMBs','Need credible social presence.','Starter reel, design or copy gig.','Low-friction entry and strong referral potential.'),
('Marketing agencies','Overflow, white-label delivery and specialist gaps.','Agency pods and priority matching.','Capacity can be sold immediately to clients.'),
('Creator-led businesses','UGC, podcast clips and repurposing.','Creator content repurposing pack.','High repeat frequency and visible proof-of-work.')])
doc.add_heading('Positioning',2)
para('For Indian businesses that need creative work shipped this week, Quick GIGS is the AI-assisted, escrow-protected marketplace that turns a rough request into a verified match and an accountable delivery—not a list of profiles to manually compare.')
doc.add_heading('Competitive frame',2)
table(['Alternative','Strength','Quick GIGS response'],[
('Open freelance marketplaces','Large supply and broad choice.','Reduce search with structured briefs, explainable match and managed quality.'),
('Traditional agencies','Accountability and strategy.','Offer agency-like coordination for smaller, faster jobs; add pods for larger accounts.'),
('In-house hiring','Control and continuity.','Serve overflow and specialist work without fixed headcount.'),
('Informal WhatsApp referrals','Trust and convenience.','Add verification, scope clarity, escrow and records.')])

# business model
page(); doc.add_heading('4. Business model and unit economics',1)
para('The default model is customer-paid service pricing plus GST. Creators/freelancers receive 100% of the service subtotal. There is no creator platform fee. A platform fee may remain configurable for a future phase, but it should stay disabled and absent from customer-facing pricing until deliberately activated.')
doc.add_heading('Revenue streams',2)
table(['Stream','How it works','Launch posture'],[
('Transaction margin / platform fee','Optional configurable fee on customer side in a future phase.','Keep at 0% while trust and repeat usage are established.'),
('Business plans','Monthly allowance, team seats, SLA and consolidated billing.','Pilot with agencies and 5–20-person marketing teams.'),
('Priority / managed pods','Higher-touch matching, account management and quality reporting.','Sell only after repeat demand is proven.'),
('Add-on services','Repurposing packs, rush delivery, creative QA or reporting.','Use as packaged, transparent line items.')])
doc.add_heading('Pricing architecture',2)
for x in ['Starter: low-risk one-off gigs for first purchase and local SMBs.','Pro: recurring reels, UGC, design or copy packs with faster turnaround.','Studio / Pod: multi-specialist campaign packs and monthly capacity.','Monthly: separate from one-time gigs; one invoice, defined allowance, overage at listed rates.','Every customer breakdown: service amount + GST. No hidden or surprise platform line item.']: bullet(x)
doc.add_heading('Illustrative transaction model',2)
table(['Metric','Illustrative assumption','Comment'],[
('Average service subtotal','₹2,499','Replace with observed blended AOV.'),
('GST shown to customer','18% = ₹450','Tax treatment must be reviewed by finance/tax adviser.'),
('Customer pays','₹2,949','Subtotal + GST only.'),
('Creator receives','₹2,499','100% of service subtotal.'),
('Platform gross contribution','₹0 initially','Monetise through future fee, plans or managed services—not creator deductions.')])

# GTM
h=doc.add_heading('5. Go-to-market strategy',1)
doc.add_heading('GTM thesis',2)
para('Start with a narrow, proof-led motion: sell a small set of repeatable outcomes to businesses that already feel the pain weekly, manually recruit supply around those outcomes, publish delivered-work proof, then expand categories and paid acquisition only after repeat and quality metrics are stable.')
doc.add_heading('Launch sequence',2)
table(['Phase','Timing','Focus','Exit criteria'],[
('0. Instrument','Weeks 1–2','Analytics, CRM, lead source, event taxonomy, creator quality rubric and service packaging.','Every lead and order has source, category, status and outcome.'),
('1. Concierge pilot','Weeks 3–8','Founder-led outreach to Lucknow/NCR agencies, D2C brands and creators. Manually guarantee matching.','30 paid orders, ≥80% on-time, 10 repeat buyers.'),
('2. Proof engine','Months 3–4','Publish case studies, delivered-work marquee, before/after briefs and customer testimonials.','20 usable proof assets and repeatable case-study format.'),
('3. Repeatable acquisition','Months 5–8','Partnerships, SEO/AEO, webinars, referral loops and controlled paid search/social.','CAC payback and second-order rate meet thresholds.'),
('4. Scale pods','Months 9–12','Agency white-label, monthly plans, account-based outbound and city expansion.','Three profitable acquisition channels and reliable supply density.')])

doc.add_heading('Channel plan',2)
table(['Channel','Audience','Offer / creative','Weekly operating cadence','Primary KPI'],[
('Founder-led outbound','Agency owners, D2C founders, marketing heads.','Free brief teardown + 3-specialist shortlist.','50 targeted accounts; 10 personalised demos.','Qualified meeting → first order'),
('Referral programme','Existing buyers and freelancers.','Buyer credit or priority match; never reduce creator payout.','Ask after approval and positive rating.','Referral share of new orders'),
('Partnerships','Coworking spaces, accelerators, Shopify/marketing consultants.','Member starter pack and office hours.','5 partner conversations; 1 activation.','Orders per partner'),
('SEO / AEO','High-intent searchers.','Pages for “reels editor India,” “UGC video pricing,” GST/escrow FAQs.','Publish 2 useful pages + refresh one.','Organic qualified leads'),
('Social proof','Founders, creators and marketers.','Delivered work, turnaround stories, transparent maths.','3 short posts; 1 case study.','Engaged qualified traffic'),
('Paid acquisition','Only validated segments.','Search around urgent outcomes, not generic “freelancer.”','Small tests with strict stop rules.','CAC and payback')])

# sales
h=doc.add_heading('6. Sales motion and funnel',1)
doc.add_heading('Buyer funnel',2)
table(['Step','Definition','Target benchmark after learning'],[
('Reach','Qualified target account or intent visitor.','Track by segment and source.'),
('Lead','Brief, enquiry or contact with business identity.','Landing → lead: 3–8% depending on channel.'),
('Qualified lead','Clear use case, budget/timing and reachable contact.','≥50% of captured leads.'),
('First order','Paid service order placed.','Qualified lead → order: 15–30% in concierge phase.'),
('Approved delivery','Work accepted and payout released.','≥90% on-time; ≥4.6 rating.'),
('Repeat order','Second order within 90 days.','35% target by month 12.')])
doc.add_heading('Core sales scripts',2)
for x in ['“Send me the rough request you would normally put in WhatsApp. I will turn it into a clear brief and show you who could start.”','“You do not need to hire a team for the next campaign. Start with one deliverable, pay GST transparently, and approve before the freelancer is paid.”','“For agencies: keep the client relationship; use a white-label pod for overflow, with a shared delivery record and one consolidated invoice.”']: bullet(x)
doc.add_heading('CRM fields required',2)
para('Source, segment, city, category, monthly demand, urgency, budget band, first response time, assigned owner, match shown, order status, delivery status, rating, repeat date, reason lost and next action.')

# supply
h=doc.add_heading('7. Supply-side growth and quality',1)
table(['Area','Policy / playbook'],[
('Recruitment','Invite specialists through LinkedIn, Instagram, creator communities, referrals and agency networks; prioritise proof of work over follower count.'),
('Verification','Identity/contact checks, portfolio review, skill category, availability, rate, turnaround and sample brief response.'),
('Agency teams','Allow multi-specialist companies to join as one team profile with named capabilities and a single commercial contact.'),
('Activation','First approved portfolio, response within SLA, availability calendar and acceptance of payout terms.'),
('Quality','Category-specific checklist, brief completeness score, delivery punctuality, revision rate and buyer rating.'),
('Payout trust','Creator receives 100% of service subtotal; show expected amount before accepting; release after approval or policy timeout.')])

# metrics
page(); doc.add_heading('8. Metrics, targets and operating dashboard',1)
doc.add_heading('North-star metric',2)
para('Approved, repeatable work delivered: the number of customer orders approved on time with a positive rating and a clear creator payout. This prevents growth that is only lead volume or unapproved GMV.')
table(['Metric','Formula','Owner','Green threshold'],[
('Qualified lead rate','Qualified leads / captured leads','Growth + sales','≥50%'),
('First-order conversion','First orders / qualified leads','Sales','≥20% pilot'),
('Time to first match','Median minutes from brief to shortlist','Product + supply','<15 minutes'),
('On-time delivery','On-time approved orders / approved orders','Operations','≥90%'),
('Approval rating','Average buyer rating after delivery','Operations','≥4.6/5'),
('Repeat rate','Buyers with second order within 90 days / first-order buyers','Lifecycle','≥35%'),
('Creator activation','Creators with one approved order / verified creators','Supply','≥50%'),
('CAC payback','CAC / gross contribution per buyer per month','Finance','Within 3 months when paid scale begins')])
doc.add_heading('Instrumentation checklist',2)
for x in ['UTM source/medium/campaign on every acquisition link.','Events: brief_started, brief_completed, lead_submitted, match_viewed, order_created, delivery_submitted, order_approved, payout_released, repeat_order.','Dashboard by segment, category, city, source and cohort—not only totals.','Weekly customer interviews: 5 buyers and 5 creators in the pilot.','Monthly cohort review: activation, approval, rating, repeat and contribution.']: bullet(x)

# budget
h=doc.add_heading('9. First-year operating plan and budget',1)
para('Use a staged budget. Do not commit to large paid media before the concierge pilot proves conversion, fulfilment and repeat behavior.')
table(['Workstream','Months 1–3','Months 4–6','Months 7–12','Control'],[
('Founder / sales','High','High','Medium','Weekly pipeline review'),
('Content and proof','Medium','High','High','Asset-level lead attribution'),
('Partnerships','Low','Medium','High','Orders per partner'),
('Paid acquisition','Test only','Small tests','Scale winners only','Stop at CAC ceiling'),
('Operations / QA','High','High','High','On-time and rating dashboard'),
('Product / automation','Instrumentation','Repeat workflows','Pods and billing','Impact per shipped feature')])
doc.add_heading('Budget guardrails',2)
for x in ['Set a fixed monthly test budget for each paid channel; pause any ad set after a pre-agreed spend without a qualified lead.','Pay for proof assets and customer outcomes before broad brand campaigns.','Treat founder time as a real cost in channel comparisons.','Keep GST, refunds, payment processing, support and creator payouts separate in the ledger.','No creator acquisition incentive should reduce the creator’s 100% service-subtotal payout.']: bullet(x)

# risks
h=doc.add_heading('10. Risks and mitigations',1)
table(['Risk','Early signal','Mitigation'],[
('Low trust / marketplace cold start','Visitors browse but do not submit briefs.','Concierge matching, transparent GST, escrow explanation, proof and human callback.'),
('Inconsistent quality','High revision rate or low ratings.','Narrow categories, verified samples, QA checklist and pause weak supply.'),
('Supply fragmentation','Creators sign up but remain inactive.','Recruit against real demand, show expected payout, route first job quickly.'),
('Paid acquisition too early','Leads cost more than contribution.','Pilot first, segment by use case, enforce CAC stop rules.'),
('Scope creep','Orders require repeated revisions.','Structured briefs, deliverables, revision policy and change-request pricing.'),
('Tax / compliance complexity','Invoice or GST exceptions.','Have a qualified Indian tax adviser review GST, TDS, invoices, contracts and payout records.')])

# 90 day
h=doc.add_heading('11. 90-day action plan',1)
table(['Week','Actions','Deliverable'],[
('1–2','Choose 3 beachhead packages; instrument funnel; define QA and creator verification rubric.','Dashboard, package pages, scripts and scorecards.'),
('3–4','Recruit 30–40 supply-side specialists; contact 100 target businesses; run founder demos.','First 10 paid orders and baseline conversion.'),
('5–8','Deliver manually; interview every buyer; publish first proof assets; launch referral ask.','30 paid orders, 10 repeat opportunities, 5 case studies.'),
('9–10','Improve onboarding, matching explanations, pricing pages and lifecycle messages.','Higher brief completion and faster match time.'),
('11–12','Test one partnership and one paid-intent campaign; review cohort economics.','Scale / stop decision for each channel.')])

# conclusion
h=doc.add_heading('12. Decision framework',1)
para('Quick GIGS should scale only when three things are true at the same time: customers return, creators are paid fairly and on time, and each acquisition channel can be measured to an economically sensible outcome. The first growth milestone is not maximum signups; it is a reliable loop of brief → verified match → approved delivery → repeat order.')
para('Recommended immediate decision: operate a narrow Lucknow/NCR and remote-India concierge pilot around recurring video, UGC, design and copy packages. Use the resulting proof to expand into agencies, monthly pods and paid acquisition. Keep the customer promise simple, keep the creator payout policy explicit, and make every number visible to the operator.')

# save
# keep tables from splitting rows when possible
for t in doc.tables:
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)
doc.save(OUT)
print(OUT)
